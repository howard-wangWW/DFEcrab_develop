# -*- coding: utf-8 -*-
"""
FileService - 会话附件解析与注入（批次12：文件实体化生命周期）

对齐 LibreChat「先上传原文件，再本地解析，解析文本才进上下文」：
  - 上传即落盘 + 解析一次，摘要缓存进注册表（registry.py）
  - 对话按会话自动注入 active 附件摘要（summary_cache 直取，不重解析）
  - 二进制不进上下文；原文件保留供下载/追溯
  - × 删除 / 删会话 → 由 file_handler / session 联动清 registry + 落盘

阶段2刻意不再支持请求体内联 base64（阶段1 兼容已移除）。
txt/md/log/csv 用标准库；xlsx 需 openpyxl（缺失友好降级）。
"""
from __future__ import annotations

import json
import logging
import re
from pathlib import Path
from typing import Any, Dict, List, Optional

logger = logging.getLogger(__name__)

_PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent

# ★ 文件上传放开（2026-09-22）：扩展名分组
#   _TEXT_LIKE_EXTS：按文本解析（编码探测 + 预算截断，与 .txt 同质量）
#   _IMAGE_EXTS    ：按图片解析（OCR 提取文字；引擎不可用时明确说明，不产出乱码）
_TEXT_LIKE_EXTS = (
    ".txt", ".md", ".markdown", ".log",
    ".json", ".xml", ".yaml", ".yml", ".toml",
    ".ini", ".conf", ".cfg",
    ".py", ".js", ".ts", ".sh", ".sql", ".html", ".htm", ".css",
    ".tsv",
)

_IMAGE_EXTS = (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp")


def _decode_text(raw: bytes) -> str:
    """编码探测：utf-8 → gbk → utf-16 → 兜底 replace（中文现场常见 GBK/UTF-8 混用）。"""
    for enc in ("utf-8", "gbk", "utf-16"):
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, ValueError):
            continue
    return raw.decode("utf-8", errors="replace")


# ────────────────────────────── 解析器（上传时执行一次）──────────────────────────────

def _parse_text_file(raw: bytes, max_chars: int) -> Dict[str, Any]:
    """txt/md/log：编码探测 + 受控截断（头尾保留）。"""
    text = _decode_text(raw)
    if len(text) <= max_chars:
        return {"type": "text", "chars": len(text), "content": text}
    head = text[: int(max_chars * 0.7)]
    tail = text[-int(max_chars * 0.25):]
    return {
        "type": "text",
        "chars": len(text),
        "content": f"{head}\n\n…[内容过长已截断，原 {len(text)} 字符，省略中部]…\n\n{tail}",
    }


def _parse_csv_file(raw: bytes, summary_rows: int) -> Dict[str, Any]:
    """csv：列名 + 类型 + 总行数 + 前 N 行（标准库，无 pandas）。"""
    import csv
    import io

    text = _decode_text(raw)
    reader = csv.reader(io.StringIO(text))
    try:
        rows = list(reader)
    except Exception as e:
        return {"type": "csv", "error": f"csv 解析失败: {e}", "total_rows": 0, "columns": []}
    if not rows:
        return {"type": "csv", "error": "空文件", "total_rows": 0, "columns": []}
    header = [c.strip() for c in (rows[0] or [])]
    data_rows = rows[1:]
    total_rows = len(data_rows)
    preview = data_rows[:summary_rows]
    col_types: List[str] = []
    if header:
        for ci in range(len(header)):
            _sample = ""
            for r in preview:
                if ci < len(r) and r[ci].strip():
                    _sample = r[ci].strip()
                    break
            if not _sample:
                col_types.append("?")
            elif re.fullmatch(r"[-+]?\d+(\.\d+)?([eE][-+]?\d+)?", _sample):
                col_types.append("number")
            else:
                col_types.append("text")
    return {
        "type": "csv",
        "columns": list(zip(header, col_types)),
        "total_rows": total_rows,
        "preview": preview,
        "truncated": total_rows > len(preview),
    }


def _parse_xlsx_file(path: Path, summary_rows: int, max_sheets: int) -> Dict[str, Any]:
    """xlsx：逐 sheet 摘要（openpyxl；缺失时友好降级）。"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {"type": "xlsx", "error": "缺少 openpyxl 依赖，无法解析 xlsx（请 pip install openpyxl）"}
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        return {"type": "xlsx", "error": f"xlsx 打开失败: {e}"}
    sheets = []
    for sheet in wb.sheetnames[:max_sheets]:
        ws = wb[sheet]
        sheet_rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > summary_rows:
                break
            sheet_rows.append(["" if c is None else str(c) for c in row])
        sheets.append({"name": sheet, "preview": sheet_rows})
    return {"type": "xlsx", "sheets": sheets, "sheet_count": len(wb.sheetnames)}


def _parse_image_file(raw: bytes, max_chars: int, image_ocr: bool = True) -> Dict[str, Any]:
    """png/jpg/...：OCR 提取图片文字（离线、进程内引擎；不可用时明确降级）。

    OCR 引擎复用知识库同一实现（src/knowledge/ocr/engine.py，进程级单例）：
    缺少后端时 is_available() 为 False —— 此时给出明确说明，而不是把二进制
    当文本解码成乱码塞进上下文（反幻觉：宁可说"没识别到"，不可编造内容）。
    """
    info: Dict[str, Any] = {"type": "image", "chars": 0, "content": ""}
    try:
        from src.knowledge.ocr.engine import get_ocr_engine, image_pixel_size
    except Exception as e:
        info["error"] = f"图片已上传，但 OCR 模块不可用（导入失败: {e}）"
        return info

    size = image_pixel_size(raw)
    if size:
        info["width"], info["height"] = int(size[0]), int(size[1])

    if not image_ocr:
        info["error"] = "图片已上传，但 image_ocr=false，未做文字提取"
        return info

    engine = get_ocr_engine()
    if not engine.is_available():
        info["error"] = ("图片已上传，但 OCR 引擎不可用，未提取文字"
                         f"（{engine.error or '未安装 OCR 后端'}）")
        return info

    text = (engine.image_to_text(raw) or "").strip()
    if not text:
        info["error"] = "图片已上传，OCR 未识别到文字"
        return info

    info["chars"] = len(text)
    if len(text) > max_chars:
        head = text[: int(max_chars * 0.7)]
        tail = text[-int(max_chars * 0.25):]
        text = f"{head}\n…[OCR 文本过长已截断，原 {info['chars']} 字符]…\n{tail}"
    info["content"] = text
    return info


def parse_file_by_ext(filename: str, raw: bytes, path: Optional[Path] = None,
                      max_text_chars: int = 12000, summary_rows: int = 10,
                      max_xlsx_sheets: int = 10,
                      image_ocr: bool = True) -> Dict[str, Any]:
    """按扩展名解析原始字节 → 解析结果 dict（供生成 summary_cache）。

    ★ 文件上传放开（2026-09-22）：文本类扩展名统一走与 .txt 同质量的解析
      （编码探测 + 预算截断），图片走 OCR；不再落进"8000 字符裸文本"兜底。
    """
    ext = Path(filename).suffix.lower()
    if ext == ".csv":
        return _parse_csv_file(raw, summary_rows)
    if ext == ".xlsx":
        return _parse_xlsx_file(path or Path(filename), summary_rows, max_xlsx_sheets)
    if ext in _TEXT_LIKE_EXTS:
        return _parse_text_file(raw, max_text_chars)
    if ext in _IMAGE_EXTS:
        return _parse_image_file(raw, max_text_chars, image_ocr=image_ocr)
    # 兜底：按文本尝试
    return {"type": "text", "chars": len(raw), "content": _decode_text(raw)[:8000]}


def render_summary(filename: str, parsed: Dict[str, Any], summary_rows: int = 10) -> str:
    """解析结果 → 注入用的单文件摘要文本（存 summary_cache）。"""
    lines = [f"[附件] {filename}"]
    if parsed.get("error"):
        lines.append(f"（{parsed['error']}）")
        return "\n".join(lines)
    t = parsed.get("type")
    if t == "text":
        lines.append(f"（文本 · {parsed['chars']} 字符）")
        lines.append(parsed["content"])
    elif t == "image":
        _dims = ""
        if parsed.get("width") and parsed.get("height"):
            _dims = f" {parsed['width']}×{parsed['height']}px"
        lines.append(f"（图片{_dims} · OCR {parsed.get('chars', 0)} 字符）")
        if parsed.get("content"):
            lines.append(parsed["content"])
    elif t == "csv":
        lines.append(f"（CSV · {parsed['total_rows']} 数据行 × {len(parsed['columns'])} 列）")
        cols = " | ".join(f"{c}({ct})" for c, ct in parsed["columns"])
        lines.append(f"列: {cols}")
        lines.append(f"前 {len(parsed['preview'])} 行:")
        for r in parsed["preview"]:
            lines.append(" | ".join(r))
        if parsed.get("truncated"):
            lines.append(f"…（共 {parsed['total_rows']} 行，仅展示前 {len(parsed['preview'])} 行）")
    elif t == "xlsx":
        lines.append(f"（Excel · {parsed['sheet_count']} 个 sheet，仅展示前 {len(parsed['sheets'])} 个）")
        for sh in parsed["sheets"]:
            lines.append(f"## Sheet: {sh['name']}")
            for r in sh["preview"]:
                lines.append(" | ".join(r))
    return "\n".join(lines)


# ────────────────────────────── FileService（上传 / 注入）──────────────────────────────

class FileService:
    """会话附件：上传入库 + 按会话注入上下文。"""

    @staticmethod
    def _attachment_meta(record: Dict[str, Any]) -> Dict[str, Any]:
        filename = str(record.get("filename") or "")
        return {
            "file_id": record.get("file_id", ""),
            "filename": filename,
            "size": int(record.get("size") or 0),
            "type": record.get("type") or Path(filename).suffix.lower(),
            "status": record.get("status", "active"),
            "created_at": record.get("created_at", ""),
        }

    @staticmethod
    def store_upload(session_id: str, user_id: str, filename: str,
                     raw: bytes, cfg: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
        """上传：校验 → 落盘 → 解析 → 摘要缓存 → 注册。

        Returns:
            registry 记录（含 file_id / summary_cache）
        """
        from .registry import session_storage_dir, register
        cfg = cfg or FileService._cfg()
        safe_name = "".join(c for c in Path(filename).name if c.isalnum() or c in "._- ")[:120].strip()
        if not safe_name:
            safe_name = "file"
        sess_dir = session_storage_dir(session_id)
        sess_dir.mkdir(parents=True, exist_ok=True)
        from .registry import new_file_id
        fid = new_file_id()
        storage = str(sess_dir / f"{fid}_{safe_name}")
        Path(storage).write_bytes(raw)

        parsed = parse_file_by_ext(
            safe_name, raw, Path(storage),
            max_text_chars=cfg.get("max_text_chars", 12000),
            summary_rows=int(cfg.get("summary_rows", 10)),
            max_xlsx_sheets=int(cfg.get("max_xlsx_sheets", 10)),
            image_ocr=bool(cfg.get("image_ocr", True)),
        )
        summary = render_summary(safe_name, parsed, int(cfg.get("summary_rows", 10)))
        return register({
            "session_id": session_id,
            "user_id": user_id,
            "filename": safe_name,
            "size": len(raw),
            "type": Path(safe_name).suffix.lower(),
            "storage": storage,
            "summary_cache": summary,
        })

    @staticmethod
    def claim_new_attachments(session_id: str, message_id: str, user_id: str,
                              file_ids: Optional[List[str]] = None) -> List[Dict[str, Any]]:
        """将本轮新附件绑定到一条 user 消息，并返回消息侧快照。

        规则：
          - 显式 file_ids：只认领这些 active 且属于本会话/当前用户的附件
          - 未传 file_ids：自动认领当前会话里尚未绑定 message_id 的 active 附件
          - 每个附件只绑定一次；已绑定的附件不会重复挂到后续消息
        """
        from .registry import attach_to_message, get, list_unclaimed

        claimed: List[Dict[str, Any]] = []
        if file_ids:
            candidates: List[Dict[str, Any]] = []
            seen = set()
            for file_id in file_ids:
                fid = str(file_id or "").strip()
                if not fid or fid in seen:
                    continue
                seen.add(fid)
                rec = get(fid)
                if rec is None:
                    continue
                if rec.get("session_id") != session_id or rec.get("user_id") != user_id:
                    continue
                if rec.get("status") != "active":
                    continue
                if rec.get("message_id"):
                    continue
                candidates.append(rec)
        else:
            candidates = list_unclaimed(session_id, user_id=user_id)

        for rec in candidates:
            bound = attach_to_message(str(rec.get("file_id") or ""), message_id)
            if bound is None:
                continue
            claimed.append(FileService._attachment_meta(bound))
        return claimed

    @staticmethod
    def purge_session(session_id: str) -> Dict[str, Any]:
        """删会话联动清理（批次12步骤3）：该会话附件全量下线。

        幂等操作，用于会话被删除/过期时回收附件资源：
          1) 注册表：物理移除该会话全部记录（remove_by_session）
          2) 落盘：rmtree 会话专属目录 data/files/{session_id}/
          3) 注入侧无需额外动作——build_session_context 只查注册表，
             记录没了自然不再注入（叉掉即失效语义的最终形态）

        Returns:
            {"removed_records": int, "removed_path": str|None, "error": str|None}
        """
        import shutil
        from .registry import remove_by_session, session_storage_dir

        try:
            n = remove_by_session(session_id)
        except Exception as e:
            return {"removed_records": 0, "removed_path": None, "error": f"注册表清理失败: {e}"}

        removed_path = None
        try:
            d = session_storage_dir(session_id)
            if d.exists():
                shutil.rmtree(d, ignore_errors=True)
                removed_path = str(d)
        except Exception as e:
            return {
                "removed_records": n,
                "removed_path": None,
                "error": f"物理目录清理失败: {e}",
            }
        return {"removed_records": n, "removed_path": removed_path, "error": None}

    @staticmethod
    def _cfg() -> Dict[str, Any]:
        try:
            from src.config.app_config import section
            return section("file_upload") or {}
        except Exception:
            return {}

    @staticmethod
    def _inject_texts(cfg: Dict[str, Any]) -> Dict[str, str]:
        """注入块文案/后缀标签（唯一来源：config.file_upload.inject_texts；代码不写死文案）。"""
        raw = cfg.get("inject_texts") or {}
        if not isinstance(raw, dict):
            return {}
        return {str(k): str(v) for k, v in raw.items() if isinstance(v, str)}

    @staticmethod
    def _record_text(rec: Dict[str, Any], texts: Dict[str, str]) -> str:
        """单条附件记录 → 注入文本（仅 active；无内容返回 ""，不做任何猜测/降级）。"""
        if not rec or rec.get("status") != "active":
            return ""
        return str(rec.get("summary_cache")
                   or f"{rec.get('filename', '')} {texts.get('no_summary', '')}".strip())

    @staticmethod
    def _join_texts(records: List[Dict[str, Any]], texts: Dict[str, str]) -> str:
        """多条附件记录 → 注入块（过滤空项，保持注册表顺序）。"""
        return "\n\n".join(
            t for t in (FileService._record_text(r, texts) for r in records) if t
        )

    @staticmethod
    def build_message_context(session_id: str, message_id: str = "",
                              context_length: int = 32768) -> str:
        """本轮消息的附件注入块（附件属于其所属消息，紧随用户问题 → 指代由位置消解）。

        规则（数据 + 配置驱动，零关键词、零类型名单、零兜底猜测）：
          - 只取 registry 中绑定到 message_id 的 active 附件（message_id 为空即无"本轮"）
          - 预算 = context_length × budget_ratio（config）
          - inject_scope=session：回退旧行为（会话全部 active 附件摘要）

        Returns:
            str 注入块；无附件返回 ""
        """
        from .registry import list_by_message
        cfg = FileService._cfg()
        if str(cfg.get("inject_scope") or "").strip().lower() == "session":
            return FileService.build_session_context(session_id, context_length)

        current = list_by_message(message_id, session_id=session_id)
        if not current:
            return ""

        texts = FileService._inject_texts(cfg)
        ratio = cfg.get("budget_ratio")
        budget = int(context_length * float(ratio)) if ratio else 0
        body = FileService._join_texts(current, texts)
        if budget and len(body) > budget:
            body = body[:budget] + "\n" + texts.get("truncated", "")
        logger.info(
            f"[Pipeline] 附件注入 本轮附件={len(current)} 字符={len(body)} "
            f"(session={session_id})"
        )
        return body

    @staticmethod
    def build_history_message_context(session_id: str, message: Dict[str, Any]) -> str:
        """某条历史消息的附件注入块（附件随其所属消息进入上下文）。

        对齐成熟平台：附件不属于"会话"，只属于它挂载的那条消息 —— 因此追问能否看到
        附件内容，取决于该条历史消息是否仍在上下文窗口内（滚出即消失）。

        归属（数据驱动，取并集去重，不猜测）：
          - registry 中 message_id == 该消息 id 的附件
          - 该消息 attachments[].file_id
        仅注入 status=active 的附件：DELETE 后立即失效（以注册表当前状态为准）。

        inject_scope=session（回退模式）时返回 ""：该模式下会话全量摘要已由
        build_message_context → build_session_context 统一注入，此处不得再注入（防重复）。

        Returns:
            str 注入块；无附件返回 ""
        """
        from .registry import get, list_by_message
        cfg = FileService._cfg()
        if str(cfg.get("inject_scope") or "").strip().lower() == "session":
            return ""

        mid = str(message.get("id") or "")
        ids: List[str] = []
        if mid:
            ids.extend(str(r.get("file_id") or "")
                       for r in list_by_message(mid, session_id=session_id))
        for item in (message.get("attachments") or []):
            if isinstance(item, dict) and item.get("file_id"):
                ids.append(str(item["file_id"]))

        records: List[Dict[str, Any]] = []
        seen = set()
        for fid in ids:
            if not fid or fid in seen:
                continue
            seen.add(fid)
            rec = get(fid)
            if rec:
                records.append(rec)
        return FileService._join_texts(records, FileService._inject_texts(cfg))

    @staticmethod
    def build_session_context(session_id: str, context_length: int = 32768) -> str:
        """回退路径（inject_scope=session）：会话全部 active 附件摘要。

        预算 = context_length × budget_ratio、文案取自 inject_texts —— 与 message 模式同一套
        配置，不写死文案、不写死预算下限。

        Returns:
            str 注入块；无附件返回 ""
        """
        from .registry import list_by_session
        cfg = FileService._cfg()
        texts = FileService._inject_texts(cfg)
        files = list_by_session(session_id)
        if not files:
            return ""
        ratio = cfg.get("budget_ratio")
        budget = int(context_length * float(ratio)) if ratio else 0
        block = FileService._join_texts(files, texts)
        if budget and len(block) > budget:
            block = block[:budget] + "\n" + texts.get("truncated", "")
        return block
